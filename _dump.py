import openpyxl, io

out = open(r'd:/CODE/LPS/_excel_dump.txt', 'w', encoding='utf-8')

wb = openpyxl.load_workbook(r'Database/Scripts/APS/订单表APS用字段对照(1).xlsx', data_only=True)

for ws in wb.worksheets:
    out.write(f"\n\n################## SHEET: {ws.title}  (rows={ws.max_row}, cols={ws.max_column}) ##################\n")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        cells = ['' if c is None else str(c) for c in row]
        if any(c.strip() for c in cells):
            out.write(f"[{i}] " + "\t".join(cells) + "\n")

out.close()
print("wrote _excel_dump.txt")
